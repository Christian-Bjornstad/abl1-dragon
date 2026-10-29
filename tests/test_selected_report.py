import os,sys,tempfile,unittest,zipfile
from pathlib import Path
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QApplication
from abl1_workstation.models import VariantModel
from abl1_workstation.domain import load_pilot
from abl1_workstation.report import export_report

class SelectedReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def test_selection_distinguishes_samples_and_survives_model_sort(self):
        v=load_pilot().data['variants'][0];other={**v,'sample':'Other'}
        model=VariantModel([v,other])
        self.assertEqual(model.data(model.index(0,0),Qt.ItemDataRole.CheckStateRole),Qt.CheckState.Unchecked)
        self.assertTrue(model.setData(model.index(0,0),Qt.CheckState.Checked,Qt.ItemDataRole.CheckStateRole))
        self.assertEqual(model.selected_rows(),[v])
        self.assertEqual(model.data(model.index(1,0),Qt.ItemDataRole.CheckStateRole),Qt.CheckState.Unchecked)
        model.replace([other]);self.assertEqual(model.selected_rows(),[])
    def test_report_contains_only_variant_sheets_with_unique_protein_names_and_images(self):
        from openpyxl import load_workbook
        from PyQt6.QtGui import QImage,QColor
        from abl1_workstation.workflow import row_key
        session=load_pilot();v=session.data['variants'][0];other={**v,'sample':'Other'}
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);png=root/'proof.png';img=QImage(60,30,QImage.Format.Format_RGB32);img.fill(QColor('gray'));img.save(str(png))
            path=root/'report.xlsx'
            results={__import__('abl1_workstation.workflow',fromlist=['allele_key']).allele_key(v):[
                {'provider':'clinvar','status':'found','summary':'TEST_ONLY','url':'https://www.ncbi.nlm.nih.gov/clinvar/','screenshot':str(png),'identity_verified':True}]}
            export_report(session,[v,other],results,path,providers=['clinvar'],igv_images={row_key(v):str(png)})
            book=load_workbook(path)
            self.assertEqual(len(book.sheetnames),2);self.assertNotIn('Resultater',book.sheetnames)
            self.assertTrue(all('Thr315Ile' in s for s in book.sheetnames));self.assertEqual(len(set(s.lower() for s in book.sheetnames)),2)
            self.assertEqual(len(book.worksheets[0]._images),2);self.assertEqual(len(book.worksheets[1]._images),1)
            text=' '.join(str(c.value or '') for row in book.worksheets[1] for c in row)
            self.assertIn('Other',text);self.assertIn('Ikke tilgjengelig',text);book.close()
            with zipfile.ZipFile(path) as z:self.assertEqual(len([n for n in z.namelist() if n.startswith('xl/media/')]),3)
    def test_empty_selection_does_not_overwrite_existing_report(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'report.xlsx';path.write_bytes(b'KEEP')
            with self.assertRaises(ValueError):export_report(load_pilot(),[],{},path)
            self.assertEqual(path.read_bytes(),b'KEEP')

if __name__=='__main__':unittest.main()

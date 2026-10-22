"""Public regression tests use only explicitly invented inputs."""
import csv
import gzip
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from abl1_workstation.domain import ReviewStore,export_csv,export_xlsx,import_reviews,is_demo,load_inputs,load_pilot,save_project,validate_package

HEADER='##fileformat=VCFv4.2\n##FORMAT=<ID=AF,Number=A,Type=Float,Description="fraction">\n##FORMAT=<ID=AD,Number=R,Type=Integer,Description="depths">\n#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tSYNTHETIC\n'

class PublicDomainTests(unittest.TestCase):
    def test_public_demo_has_no_clinical_evidence_or_private_assets(self):
        demo=load_pilot();self.assertTrue(is_demo(demo));self.assertFalse(demo.is_pilot)
        self.assertEqual([s['sample'] for s in demo.data['samples']],['DEMO-A','DEMO-B'])
        self.assertEqual(len(demo.data['variants']),3)
        for key in ('qc','read_rows','phase_rows'):self.assertEqual(demo.data[key],[])
        self.assertFalse(demo.data['provenance']['clinical_inputs_used'])
        self.assertTrue(all(v.get('phase_required') for v in demo.data['variants'] if v['sample']=='DEMO-A'))
        self.assertEqual(json.loads((ROOT/'engine/reference.json').read_text(encoding='utf-8'))['bed'],[])
        self.assertEqual(json.loads((ROOT/'app/data/igv-manifest.json').read_text())['samples'],{})
        self.assertEqual(json.loads((ROOT/'app/data/knowledge.json').read_text(encoding='utf-8'))['evidence'],[])
        self.assertFalse(list((ROOT/'app/igv').glob('*.html')))

    def test_multiallelic_import_uses_caller_values_and_never_synthesizes_phase(self):
        with tempfile.TemporaryDirectory() as temp:
            file=Path(temp)/'synthetic.vcf.gz';file.write_bytes(gzip.compress((HEADER+'chr9\t133748283\t.\tC\tT,G\t.\tPASS\t.\tGT:AF:AD:DP\t1/2:0.12,0.07:810,120,70:1000\n').encode()))
            session=load_inputs([file]);self.assertFalse(session.is_pilot)
            self.assertEqual([v['format_af'] for v in session.data['variants']],[0.12,0.07]);self.assertEqual([v['alt_ad'] for v in session.data['variants']],[120,70])
            self.assertEqual(session.data['phase_rows'],[]);self.assertEqual(session.data['qc'],[])
            self.assertNotIn('p.T315M',[v['protein_snv_alone'] for v in session.data['variants']])

    def test_assembly_ref_and_missing_fields_are_explicit(self):
        with tempfile.TemporaryDirectory() as temp:
            file=Path(temp)/'synthetic.vcf';file.write_text(HEADER+'chr9\t133748283\t.\tA\tT\t.\t.\t.\tGT:AD\t0/1:9,1\n',encoding='utf-8')
            with self.assertRaises(ValueError):load_inputs([file],assembly='GRCh38')
            row=load_inputs([file]).data['variants'][0]
            self.assertFalse(row['reference_match']);self.assertIsNone(row['protein_snv_alone']);self.assertIsNone(row['format_af']);self.assertEqual(row['filter'],'.')

    def test_project_roundtrip_and_review_conflict_preserve_local_work(self):
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp);demo=load_pilot();store=ReviewStore(folder/'live.sqlite');key=demo.variant_key(demo.data['variants'][0])
            store.put(key,{'note':'Archived synthetic note'});save_project(demo,store,folder/'saved.json')
            restored=load_inputs([folder/'saved.json']);self.assertEqual(restored.run_id,demo.run_id)
            store.put(key,{'note':'New local synthetic note'})
            self.assertEqual(import_reviews(restored,store,restored.pending_reviews),0);self.assertEqual(store.get(key)['note'],'New local synthetic note')
            fresh=ReviewStore(folder/'fresh.sqlite');self.assertEqual(import_reviews(restored,fresh,restored.pending_reviews),1)
            altered=json.loads((folder/'saved.json').read_text(encoding='utf-8'));altered['analysis']['variants'][0]['format_dp']=999
            (folder/'tampered.json').write_text(json.dumps(altered),encoding='utf-8')
            with self.assertRaises(ValueError):load_inputs([folder/'tampered.json'])

    def test_csv_and_xlsx_export_formula_text_and_current_rows(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temp:
            folder=Path(temp);demo=load_pilot();store=ReviewStore(folder/'review.sqlite');store.put(demo.variant_key(demo.data['variants'][0]),{'note':'=1+1'})
            export_csv(demo,demo.data['variants'],folder/'demo.csv',store)
            with (folder/'demo.csv').open(encoding='utf-8-sig') as stream:self.assertEqual(next(csv.DictReader(stream,delimiter=';'))['note'],"'=1+1")
            export_xlsx(demo,store,folder/'demo.xlsx');book=load_workbook(folder/'demo.xlsx',read_only=True)
            self.assertEqual(book['Varianter'].max_row,4);self.assertEqual(book['QC'].max_row,1);book.close()

    def test_imported_phase_counts_must_sum_and_recompute_protein(self):
        demo=load_pilot().data
        row={'sample':'DEMO-A','hg19_interval1':'chr9:133748282-133748284','reference_codon':'ACT','haplotype':'ATG','protein':'forged','filter':'SYNTHETIC','fragments':2,'fragments_all':2,'fragment_fraction':1,'reads':2,'reads_all':2}
        demo['phase_rows']=[row];self.assertEqual(validate_package(demo)['phase_rows'][0]['protein'],'p.Thr315Met')
        row['fragments_all']=3
        with self.assertRaises(ValueError):validate_package(demo)

if __name__=='__main__':unittest.main()

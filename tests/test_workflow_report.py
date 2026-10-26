import sys
from pathlib import Path
import tempfile
import unittest

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from abl1_workstation.domain import load_pilot
from abl1_workstation.workflow import category, allele_key, candidate_rows
from abl1_workstation.report import export_report,source_text


class WorkflowReportTests(unittest.TestCase):
    def test_empty_selection_requires_an_explicit_variant_without_claiming_negative(self):
        from openpyxl import load_workbook
        session=load_pilot()
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'empty.xlsx'
            with self.assertRaises(ValueError):export_report(session,[],{},path,providers=[],samples=['DEMO-A'])
            self.assertFalse(path.exists())
    def test_found_classification_and_evidence_are_visible_beside_record_title(self):
        text=source_text({'status':'found','classification':'Germline: Likely pathogenic',
                          'evidence_level':'criteria provided','summary':'Record title'})
        for value in ('Germline: Likely pathogenic','criteria provided','Record title'):
            self.assertIn(value,text)
    def test_candidates_keep_filters_without_promoting_artifacts(self):
        base={'chrom':'chr9','pos':133748283,'ref':'C','alt':'T','reference_match':True,
              'cdna':'c.944C>T','filter':'PASS','alt_ad':1,'sample_has_alt':True}
        self.assertEqual(category(base),'PASS')
        filtered={**base,'filter':'LowAF'}
        self.assertEqual(category(filtered),'Kontroll')
        self.assertEqual(category({**base,'reference_match':False}),'REF-avvik')
        self.assertEqual(category({**base,'normal_transcript_continuation':True}),'Transkriptfortsettelse')
        self.assertEqual(category({**filtered,'alt_ad':0}),'Øvrige')
        self.assertEqual(candidate_rows([base,filtered,{**base,'reference_match':False}]),[base,filtered])
        self.assertEqual(filtered['filter'],'LowAF')

    def test_lookup_identity_has_no_sample_and_distinguishes_alleles(self):
        v=load_pilot().data['variants'][0]
        self.assertEqual(allele_key(v),allele_key({**v,'sample':'different','chrom':'9'}))
        self.assertNotEqual(allele_key(v),allele_key({**v,'alt':'G'}))

    def test_per_variant_export_retains_phase_comments_and_source_provenance(self):
        from openpyxl import load_workbook
        session=load_pilot();rows=session.data['variants'];rows[0]['sample']='=UNTRUSTED'
        results={allele_key(rows[0]):[{'provider':'clinvar','status':'found','summary':'Test classification',
                 'checked_at':'2026-10-25T17:00Z','url':'https://www.ncbi.nlm.nih.gov/clinvar/variation/123/'}]}
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'report.xlsx'
            export_report(session,rows,results,path,providers=['clinvar','franklin'])
            book=load_workbook(path)
            self.assertEqual(len(book.sheetnames),len(rows));self.assertNotIn('Resultater',book.sheetnames)
            sheet=book.active
            self.assertEqual(sheet.cell(4,3).value,"'=UNTRUSTED")
            values=[c.value for row in sheet for c in row]
            self.assertIn('Test classification',values);self.assertIn('Ikke slått opp',values)
            self.assertIn('kodon',sheet.cell(7,3).comment.text.lower())
            comments=' '.join(c.comment.text for row in sheet for c in row if c.comment)
            self.assertIn('2026-10-25',comments);self.assertIn('/variation/123/',comments)
            self.assertEqual(sheet.freeze_panes,'C4');book.close()

from pathlib import Path
import sys,tempfile,unittest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from abl1_workstation.protein_names import protein_display,protein_key


class ProteinTests(unittest.TestCase):
    def test_display_preserves_identity_and_handles_existing_hgvs(self):
        for value,expected in [('p.V299L','p.Val299Leu'),('T315I','Thr315Ile'),
                               ('p.Thr315Ile','p.Thr315Ile'),('p.T315=','p.Thr315='),
                               ('p.W10*','p.Trp10Ter'),('p.(T315I)','p.(Thr315Ile)'),
                               ('p.M244VfsTer3','p.Met244ValfsTer3')]:
            self.assertEqual(protein_display(value),expected)
    def test_watchlist_key_only_accepts_single_known_substitution(self):
        self.assertEqual(protein_key('p.Val299Leu'),'V299L')
        self.assertEqual(protein_key('V299L'),'V299L')
        for value in ('V299L;T315I','p.(V299L)','V299*fs','NM_005157.6:p.V299L','X299L'):
            self.assertIsNone(protein_key(value))


class WatchlistTests(unittest.TestCase):
    def test_corrupt_workbook_xml_is_a_normal_import_error(self):
        from abl1_workstation.watchlist import read_watchlist
        import zipfile
        original=Path(__file__).resolve().parents[1]/'assets/ELN_2025_Tabell5.xlsx'
        with tempfile.TemporaryDirectory() as folder:
            corrupt=Path(folder)/'broken.xlsx'
            with zipfile.ZipFile(original) as source,zipfile.ZipFile(corrupt,'w') as target:
                for name in source.namelist():target.writestr(name,b'<workbook><broken' if name=='xl/workbook.xml' else source.read(name))
            with self.assertRaises(ValueError):read_watchlist(corrupt)
    def test_import_updates_model_without_selecting_rows_or_changing_internal_identity(self):
        import os
        os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
        from PyQt6.QtWidgets import QApplication
        from PyQt6.QtCore import Qt
        from abl1_workstation.domain import load_pilot,ReviewStore
        from abl1_workstation.ui import MainWindow
        app=QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as folder:
            store=ReviewStore(Path(folder)/'reviews.sqlite');window=MainWindow(store,session=load_pilot(),browser_enabled=False)
            path=Path(__file__).resolve().parents[1]/'assets/ELN_2025_Tabell5.xlsx'
            self.assertTrue(window.settings_panel.import_watchlist(path))
            self.assertFalse(window.variant_model.selected_rows())
            positive=next(i for i,v in enumerate(window.variant_model.rows) if v.get('protein_snv_alone')=='p.T315I')
            model=window.variant_model
            self.assertEqual(model.data(model.index(positive,4)),'p.Thr315Ile')
            self.assertEqual(model.data(model.index(positive,8)),'Treff')
            self.assertIsNotNone(model.data(model.index(positive,4),Qt.ItemDataRole.BackgroundRole))
            self.assertNotIn('eln_watchlist',window.session.data['variants'][positive])
            window.settings_panel._clear_watchlist()
            self.assertEqual(model.data(model.index(positive,8)),'–')
            window.close();store.close()

    def test_report_contains_source_accessions_watchlist_and_three_letter_sheet(self):
        from openpyxl import load_workbook
        from abl1_workstation.domain import load_pilot
        from abl1_workstation.report import export_report
        from abl1_workstation.workflow import allele_key
        session=load_pilot();variant=dict(session.data['variants'][0],eln_watchlist=[{'source':'ELN 2025 Table 5','scope':'CML','recommended_tki':'Ponatinib, asciminib'}])
        result={'provider':'clinvar','status':'found','transcripts':['NM_005157.6'],'proteins':['NP_005148.2'],
                'reference_hgvs':['NM_005157.6:c.944C>T','NP_005148.2:p.Thr315Ile']}
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'report.xlsx';export_report(session,[variant],{allele_key(variant):[result]},path,providers=['clinvar'])
            book=load_workbook(path)
            self.assertEqual(book.sheetnames,['Thr315Ile'])
            text=' '.join(str(c.value or '') for row in book.active for c in row)
            for expected in ('NP_005148.2','NM_005157.6:c.944C>T','ELN 2025 Table 5','CML','Ponatinib'):
                self.assertIn(expected,text)
            book.close()
    def test_supplied_table_has_29_unique_variants_and_both_amino_acid_forms_match(self):
        from abl1_workstation.watchlist import read_watchlist,matches
        rows=read_watchlist(Path(__file__).resolve().parents[1]/'assets/ELN_2025_Tabell5.xlsx')
        self.assertEqual(len(rows),29)
        self.assertEqual(len({r['protein'] for r in rows}),29)
        self.assertEqual({r['protein'] for r in rows},set('M244V Y253H E255K E255V V299L T315I F317L F317V F317I F317C T315A F359V F359I F359C A337V A337T L340Q A344P A433D G463D G463S P465S P465Q V468F F497L I502L I502N V506L V506M'.split()))
        import json,re
        from abl1_workstation.domain import _engine
        reference=json.loads((Path(__file__).resolve().parents[1]/'engine/reference.json').read_text())
        for row in rows:
            protein=re.fullmatch(r'([A-Z])(\d+)([A-Z])',row['protein'])
            at=reference['cds_start_t0']+3*(int(protein[2])-1)
            self.assertEqual(_engine().AA[reference['transcript_sequence'][at:at+3]],protein[1])
        self.assertEqual(matches({'protein_snv_alone':'p.T315I'},rows)[0]['protein'],'T315I')
        self.assertEqual(matches({'protein_snv_alone':'p.Thr315Ile'},rows)[0]['protein'],'T315I')
        self.assertFalse(matches({'protein_snv_alone':'p.T315A;T315I'},rows))
        self.assertFalse(matches({'protein_snv_alone':'p.T315T'},rows))
    def test_import_is_atomic_and_rejects_formula_wrong_gene_and_duplicates(self):
        from abl1_workstation.watchlist import read_watchlist,save_watchlist,load_watchlist
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as folder:
            destination=Path(folder)/'list.json'
            valid=[{'gene':'ABL1','protein':'T315I','transcript':'NM_005157.6','source':'test','recommended_tki':'','scope':'CML'}]
            save_watchlist(destination,valid)
            self.assertEqual(load_watchlist(destination),valid)
            for gene,protein in [('BRAF','T315I'),('ABL1','=1+1'),('ABL1','T315I/T315A')]:
                book=Workbook();sheet=book.active
                sheet.append(['gene','protein','transcript','source'])
                sheet.append([gene,protein,'NM_005157.6','test'])
                path=Path(folder)/'invalid.xlsx';book.save(path);book.close()
                with self.assertRaises(ValueError):read_watchlist(path)
                self.assertEqual(load_watchlist(destination),valid)

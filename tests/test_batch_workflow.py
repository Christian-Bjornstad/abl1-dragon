import os,sys,tempfile,time,unittest,threading
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QApplication
from abl1_workstation.domain import ReviewStore,load_pilot
from abl1_workstation.ui import MainWindow
from abl1_workstation.workflow import allele_key

class BatchWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.folder=Path(self.temp.name)
        self.store=ReviewStore(self.folder/'state.sqlite');self.window=MainWindow(self.store,browser_enabled=False)
        session=load_pilot();v=session.data['variants'][0];session.data['variants']=[v,{**v,'sample':'Other'}]
        self.window._apply_session(session)
    def tearDown(self):
        self.window.close();self.app.processEvents();self.store.close();self.temp.cleanup()
    def finish(self):
        deadline=time.monotonic()+5
        while (self.window.loading or self.window.thread is not None) and time.monotonic()<deadline:
            self.app.processEvents();time.sleep(.005)
        self.app.processEvents();self.assertFalse(self.window.loading)
    def test_marks_survive_filter_and_batch_queries_each_allele_once(self):
        w=self.window;self.assertFalse(w.lookup_button.isEnabled());w.mark_visible()
        self.assertEqual(len(w.variant_model.selected_rows()),2)
        w.search.setText('no visible result');self.assertEqual(len(w.variant_model.selected_rows()),2)
        with patch('abl1_workstation.lookup.LookupService.lookup',return_value=[{'provider':'clinvar','status':'not_found'}]) as lookup:
            w.lookup_selected();self.finish();self.assertEqual(lookup.call_count,1)
        self.assertEqual(len(w.lookup_results),1)
        w.clear_marks();self.assertFalse(w.excel_button.isEnabled())
    def test_export_captures_each_sample_instead_of_reusing_current_variant(self):
        from PyQt6.QtGui import QImage,QColor
        from openpyxl import load_workbook
        w=self.window;w.mark_visible();w.settings.reference='synthetic_reference'
        w.run_files=SimpleNamespace(alignments={s:SimpleNamespace(sample=s) for s in ('DEMO-A','Other')})
        viewer=w.read_viewer;captured=[]
        def show(alignment,reference,variant):
            viewer.current_alignment=alignment;viewer.save_image_button.setEnabled(True);captured.append((alignment.sample,variant['sample']))
        def capture(path,callback):
            image=QImage(160,100,QImage.Format.Format_RGB32);image.fill(QColor('gray'));image.save(str(path));callback(str(path),None)
        path=self.folder/'exports'
        with patch.object(viewer,'show_alignment',side_effect=show),patch.object(viewer,'capture_png',side_effect=capture):
            w.start_export(path);self.finish()
        self.assertEqual(captured,[('DEMO-A','DEMO-A'),('Other','Other')])
        books=[load_workbook(p) for p in sorted(path.glob('*.xlsx'))];self.assertEqual(len(books),2)
        self.assertEqual([len(book.worksheets[0]._images) for book in books],[1,1])
        for book in books:self.assertEqual(book.sheetnames,['Thr315Ile']);book.close()
        self.assertEqual(w.pages.currentIndex(),0)
    def test_cancelled_export_never_creates_a_partial_workbook(self):
        w=self.window;w.mark_visible();path=self.folder/'cancelled'
        w.start_export(path);w.cancel_export();self.app.processEvents()
        self.assertFalse(path.exists());self.assertIsNone(w._export);self.assertFalse(w.loading)

    def test_first_source_is_visible_and_saved_while_next_source_is_still_running(self):
        from abl1_workstation.lookup import canonical_identity
        from abl1_workstation.evidence_store import EvidenceStore
        w=self.window;w.mark_visible();release=threading.Event();started=threading.Event()
        variant=w.variant_model.selected_rows()[0]
        def lookup(v,providers,**options):
            item={'provider':'clinvar','status':'found','summary':'Synthetic persisted evidence',
                  'identity':canonical_identity(v),'identity_verified':True}
            options['completed'](item);started.set();release.wait(4)
            return [item]
        try:
            with patch('abl1_workstation.lookup.LookupService.lookup',side_effect=lookup):
                w.lookup_selected();self.assertTrue(started.wait(2))
                deadline=time.monotonic()+2
                while not w.lookup_results and time.monotonic()<deadline:self.app.processEvents();time.sleep(.005)
                self.assertTrue(w.loading)
                self.assertEqual(w.lookup_results[allele_key(variant)][0]['summary'],'Synthetic persisted evidence')
                self.assertEqual(EvidenceStore(self.folder).load(variant)[0]['provider'],'clinvar')
                release.set();self.finish()
            w._apply_session(load_pilot())
            self.assertEqual(w.lookup_results[allele_key(variant)][0]['summary'],'Synthetic persisted evidence')
        finally:release.set();self.finish()

    def test_partial_capture_is_visible_as_partial_even_when_one_picture_is_available(self):
        w=self.window;variant=w.current_variant
        w._lookup_finished((allele_key(variant),[{'provider':'clinvar','status':'found','screenshot':'synthetic.png','screenshot_error':'A second view is missing'}]))
        self.assertIn('delvis',w.source_badges['clinvar'].text().lower())
        self.assertNotIn('bilde klart',w.source_badges['clinvar'].text().lower())

if __name__=='__main__':unittest.main()

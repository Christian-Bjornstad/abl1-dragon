import os,sys,tempfile,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from openpyxl import load_workbook
from abl1_workstation.domain import load_pilot
from abl1_workstation.report import export_sample_reports,sample_report_paths
from abl1_workstation.workflow import allele_key


class SampleReportTests(unittest.TestCase):
    def test_two_samples_produce_separate_workbooks_with_only_their_marked_variants(self):
        session=load_pilot();variant=session.data['variants'][0]
        rows=[variant,{**variant,'sample':'PRØVE-2'},{**variant,'sample':'PRØVE-2','protein_snv_alone':'p.T315M','alt':'G'}]
        with tempfile.TemporaryDirectory() as folder:
            outputs=export_sample_reports(session,rows,{},folder,providers=[])
            self.assertEqual([p.name for p in outputs],['DEMO-A_DRAGEN_Resultater.xlsx','PRØVE-2_DRAGEN_Resultater.xlsx'])
            first=load_workbook(outputs[0]);second=load_workbook(outputs[1])
            self.assertEqual(first.sheetnames,['Thr315Ile']);self.assertEqual(second.sheetnames,['Thr315Ile','Thr315Met'])
            text=' '.join(str(c.value or '') for s in first for row in s for c in row)
            self.assertNotIn('PRØVE-2',text);first.close();second.close()
    def test_unsafe_sample_names_cannot_escape_output_folder_and_collisions_are_unique(self):
        variant=load_pilot().data['variants'][0]
        samples=['../A','..\\A','CON','LPT1','trailing. ','A/B','A\\B','a/b','Å'*300]
        with tempfile.TemporaryDirectory() as folder:
            paths=sample_report_paths([{**variant,'sample':name} for name in samples],folder)
            self.assertEqual(len(paths),len(samples));self.assertEqual(len({p.name.casefold() for p in paths.values()}),len(samples))
            for path in paths.values():
                self.assertEqual(path.resolve().parent,Path(folder).resolve())
                self.assertTrue(path.name.endswith('_DRAGEN_Resultater.xlsx'));self.assertLessEqual(len(path.name),160)
    def test_explicit_file_path_is_rejected_for_multiple_samples_without_writing_anything(self):
        variant=load_pilot().data['variants'][0]
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaises(ValueError):export_sample_reports(load_pilot(),[variant,{**variant,'sample':'Other'}],{},Path(folder)/'combined.xlsx')
            self.assertEqual(list(Path(folder).iterdir()),[])
    def test_construction_failure_does_not_replace_any_existing_sample_report(self):
        from unittest.mock import patch
        from abl1_workstation.report import export_report
        session=load_pilot();variant=session.data['variants'][0];rows=[variant,{**variant,'sample':'Other'}]
        with tempfile.TemporaryDirectory() as folder:
            paths=sample_report_paths(rows,folder)
            for path in paths.values():path.write_bytes(b'EXISTING_REPORT')
            def build(session,rows,results,path,**options):
                if rows[0]['sample']=='Other':raise OSError('Disk error')
                return export_report(session,rows,results,path,**options)
            with patch('abl1_workstation.report.export_report',side_effect=build):
                with self.assertRaises(OSError):export_sample_reports(session,rows,{},folder,providers=[])
            self.assertTrue(all(p.read_bytes()==b'EXISTING_REPORT' for p in paths.values()))
            self.assertEqual(len(list(Path(folder).iterdir())),2)

if __name__=='__main__':unittest.main()

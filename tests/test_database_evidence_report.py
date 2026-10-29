"""Synthetic evidence only. Check the actual exported workbook and PNG media."""
import sys,tempfile,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from PIL import Image,ImageDraw
from openpyxl import load_workbook
from abl1_workstation.domain import load_pilot
from abl1_workstation.report import export_report
from abl1_workstation.workflow import allele_key


class DatabaseEvidenceReportTests(unittest.TestCase):
    def test_all_labelled_pictures_are_embedded_once_and_bad_picture_does_not_lose_report(self):
        session=load_pilot();variant=session.data['variants'][0]
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);pictures=[]
            for label in ('ACMG Classification','Oncogenic Classification'):
                path=root/(str(len(pictures))+'.png')
                image=Image.new('RGB',(600,180),'white');ImageDraw.Draw(image).text((20,20),label,fill='black');image.save(path)
                pictures.append({'label':label,'path':str(path),'captured_at':'2026-10-26T10:15:00+00:00'})
            broken=root/'missing.png'
            result={'provider':'franklin','status':'found','identity_verified':True,'summary':'Synthetic result',
                    'screenshot':pictures[0]['path'],'screenshots':pictures+[{'label':'Evidence','path':str(broken)}]}
            path=root/'report.xlsx'
            export_report(session,[variant],{allele_key(variant):[result]},path,providers=['franklin','oncokb'])
            book=load_workbook(path);sheet=book.worksheets[0]
            self.assertEqual(book.sheetnames,['Thr315Ile'])
            self.assertEqual(len(sheet._images),2)
            text='\n'.join(str(c.value or '') for row in sheet for c in row)
            self.assertIn('ACMG Classification',text);self.assertIn('Oncogenic Classification',text)
            self.assertIn('Bildet kunne ikke leses',text);self.assertIn('Ikke slått opp',text)
            book.close()

    def test_source_errors_are_explained_in_export_instead_of_only_saying_error(self):
        from abl1_workstation.report import source_text
        self.assertIn('Nettverket svarte ikke',source_text({'status':'error','summary':'Nettverket svarte ikke'}))


if __name__=='__main__':unittest.main()

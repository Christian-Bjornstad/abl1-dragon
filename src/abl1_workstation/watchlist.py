"""An explicit, per-user ABL1 protein watchlist imported from Excel."""
import json,os,tempfile,zipfile
from io import BytesIO
from pathlib import Path
from .protein_names import protein_key

FIELDS=('gene','protein','transcript','recommended_tki','source','scope')


def _validate(rows):
    if not isinstance(rows,list) or not 1<=len(rows)<=2000:raise ValueError('Variantlisten må ha 1–2000 varianter.')
    result=[];seen=set()
    for row in rows:
        if not isinstance(row,dict):raise ValueError('Ugyldig variantliste.')
        item={key:str(row.get(key) or '').strip() for key in FIELDS}
        if any(len(value)>2000 or any(ord(c)<32 for c in value) for value in item.values()):raise ValueError('Ugyldig tekst i variantlisten.')
        key=protein_key(item['protein'])
        if item['gene']!='ABL1' or not key or item['transcript']!='NM_005157.6' or not item['source']:
            raise ValueError('Listen krever ABL1, enkeltvariant i protein, NM_005157.6 og kilde (source).')
        if key in seen:raise ValueError('Duplisert proteinvariant: '+key)
        seen.add(key);item['protein']=key;result.append(item)
    return result


def read_watchlist(path):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    path=Path(path)
    if path.stat().st_size>5_000_000:raise ValueError('Excel-listen er for stor (maks 5 MB).')
    try:
        with path.open('rb') as stream:raw=stream.read(5_000_001)
        if len(raw)>5_000_000:raise ValueError('Excel-listen er for stor (maks 5 MB).')
        # openpyxl can leave its archive open if workbook XML is malformed.
        # An immutable bounded buffer also avoids locking the user's Excel file.
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            if len(archive.infolist())>1000 or sum(i.file_size for i in archive.infolist())>25_000_000:
                raise ValueError('Excel-listen er for stor etter utpakking.')
        book=load_workbook(BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        try:
            if not book.worksheets:raise ValueError('Excel-listen har ingen regneark.')
            sheet=book.worksheets[0]
            if (sheet.max_row or 0)>2001 or (sheet.max_column or 0)>20:raise ValueError('Excel-listen har for mange rader eller kolonner.')
            cells=sheet.iter_rows();header=next(cells,())
            columns=[str(c.value or '').strip().lower() for c in header]
            if len(set(columns))!=len(columns) or not {'gene','protein','transcript','source'}.issubset(columns):
                raise ValueError('Kolonner må hete gene, protein, transcript og source. Valgfritt: recommended_tki og scope.')
            rows=[]
            for number,row in enumerate(cells,2):
                if number>2001 or len(row)>20:raise ValueError('Excel-listen har for mange rader eller kolonner.')
                if all(c.value is None for c in row):continue
                if any(c.data_type=='f' for c in row):raise ValueError('Formler er ikke tillatt i variantlisten.')
                rows.append(dict(zip(columns,[c.value for c in row])))
            return _validate(rows)
        finally:book.close()
    except (zipfile.BadZipFile,KeyError,StopIteration,IndexError,SyntaxError,InvalidFileException) as error:
        raise ValueError('Excel-listen kunne ikke leses.') from error


def save_watchlist(path,rows):
    rows=_validate(rows);path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    descriptor,temporary=tempfile.mkstemp(prefix='.watchlist-',suffix='.json',dir=path.parent)
    try:
        with os.fdopen(descriptor,'w',encoding='utf-8') as stream:
            json.dump(rows,stream,ensure_ascii=False);stream.flush();os.fsync(stream.fileno())
        os.replace(temporary,path)
    finally:Path(temporary).unlink(missing_ok=True)


def load_watchlist(path):
    path=Path(path)
    if not path.exists():return []
    if path.stat().st_size>5_000_000:raise ValueError('Lagret variantliste er for stor.')
    return _validate(json.loads(path.read_text(encoding='utf-8')))


def matches(variant,rows):
    key=protein_key(variant.get('protein_snv_alone'))
    return [dict(row) for row in rows if key and row['protein']==key]

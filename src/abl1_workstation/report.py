"""One worksheet per selected sample/allele, with local evidence images."""
from datetime import datetime
from pathlib import Path
import hashlib,math,os,re,tempfile
from urllib.parse import urlsplit
from . import APP_NAME,__version__
from .models import flags
from .workflow import allele_key,row_key
from .protein_names import protein_display

PROVIDER_NAMES={'clinvar':'ClinVar','oncokb':'OncoKB','franklin':'Franklin','mtbp':'MTBP'}
STATUS_NAMES={'not_found':'Ingen treff','login_required':'Innlogging kreves','license_required':'Tilgang kreves','error':'Feil','unsupported':'Ikke tilgjengelig','pending':'Analyse venter'}

def source_text(result):
    if not result:return 'Ikke slått opp'
    status=result.get('status')
    if status!='found':
        label=STATUS_NAMES.get(status,status or 'Ukjent status')
        return label+(': '+str(result['summary']) if result.get('summary') else '')
    parts=[]
    for field in ('classification','evidence_level','summary'):
        value=result.get(field)
        if field=='summary' and result.get('provider')=='franklin' and result.get('classification') and str(value).startswith('Franklin: identitetskontrollert portalresultat.'):continue
        if value and str(value) not in parts:parts.append(str(value))
    return ' · '.join(parts) or str(result.get('accession') or 'Treff')

def safe_text(value):
    if isinstance(value,str):
        value=''.join(c for c in value if ord(c)>=32 or c in '\n\t')
        if value.lstrip().startswith(('=','+','-','@')):value="'"+value
        return value[:32767]
    return value

def _sheet_name(variant,used):
    raw=str(protein_display(variant.get('protein_snv_alone')) or variant.get('cdna') or variant.get('pos') or 'Variant')
    base=re.sub(r'[\[\]:*?/\\\x00-\x1f]','_',raw.removeprefix('p.')).strip(" '") or 'Variant'
    name=base[:31];number=1
    while name.casefold() in used:
        suffix=' '+str(number+1);name=base[:31-len(suffix)]+suffix;number+=1
    used.add(name.casefold());return name

def _image(sheet,path,row):
    from PIL import Image as PillowImage
    from openpyxl.drawing.image import Image
    path=Path(path)
    if path.stat().st_size>20*1024*1024:raise ValueError('Bildet er for stort.')
    with PillowImage.open(path) as source:
        if source.format!='PNG' or source.width*source.height>32_000_000:raise ValueError('Ugyldig PNG-bilde.')
        source.verify()
    image=Image(str(path));scale=min(1,1000/image.width)
    image.width=round(image.width*scale);image.height=round(image.height*scale)
    sheet.add_image(image,'A'+str(row))
    return row+math.ceil(image.height/20)+2


def screenshot_records(result):
    """Keep compatibility with a single image; deduplicate by local path."""
    if not result or result.get('status')!='found' or result.get('identity_verified') is not True:return []
    records=[];seen=set()
    for record in result.get('screenshots') or []:
        if not isinstance(record,dict) or not record.get('path'):continue
        path=str(record['path'])
        if path not in seen:records.append(record);seen.add(path)
    if result.get('screenshot') and str(result['screenshot']) not in seen:
        records.insert(0,{'path':str(result['screenshot']),'label':'Variantresultat'})
    return records


def sample_report_paths(rows,destination):
    """Exactly one output per sample; never use a sample value as a path."""
    samples=list(dict.fromkeys(str(row.get('sample') or '') for row in rows))
    if not samples:raise ValueError('Merk minst én variant før eksport.')
    if any(not sample for sample in samples):raise ValueError('Prøvenavn mangler i valgt variant.')
    destination=Path(destination)
    if destination.suffix.lower()=='.xlsx':
        if len(samples)!=1:raise ValueError('Velg en mappe for separate Excel-filer per prøve.')
        return {samples[0]:destination}
    paths={};used=set();suffix='_DRAGEN_Resultater.xlsx'
    for sample in samples:
        base=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',sample).strip(' .') or 'Prøve'
        if re.fullmatch(r'(?:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?',base,re.I):base='_'+base
        base=base[:120].rstrip(' .');candidate=base+suffix
        if candidate.casefold() in used:
            candidate=base[:105]+'_'+hashlib.sha256(sample.encode()).hexdigest()[:12]+suffix
        used.add(candidate.casefold());paths[sample]=destination/candidate
    return paths


def export_sample_reports(session,rows,lookup_results,destination,**options):
    rows=list(rows);paths=sample_report_paths(rows,destination)
    # Construct all workbooks before replacing any existing report. A failure
    # while building one sample must not replace earlier samples' workbooks.
    staged=[]
    try:
        for sample,path in paths.items():
            path.parent.mkdir(parents=True,exist_ok=True)
            descriptor,temporary=tempfile.mkstemp(prefix='.DRAGEN-sample-',suffix='.xlsx',dir=path.parent);os.close(descriptor)
            staged.append((Path(temporary),path))
            sample_rows=[v for v in rows if str(v.get('sample') or '')==sample]
            export_report(session,sample_rows,lookup_results,temporary,**options)
        for temporary,path in staged:os.replace(temporary,path)
    finally:
        for temporary,_ in staged:temporary.unlink(missing_ok=True)
    return list(paths.values())

def export_report(session,rows,lookup_results,path,*,providers=('clinvar','oncokb','franklin','mtbp'),samples=None,selection='',igv_images=None,igv_errors=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.comments import Comment
    rows=list(rows)
    providers=tuple(p for p in providers if p in PROVIDER_NAMES)
    if not rows:raise ValueError('Merk minst én variant før eksport.')
    book=Workbook();book.remove(book.active);used=set()
    for variant in rows:
        sheet=book.create_sheet(_sheet_name(variant,used));sheet.sheet_view.showGridLines=False
        for column in 'ABCDEFGH':sheet.column_dimensions[column].width=18
        def block(row,text,*,heading=False):
            sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=8)
            cell=sheet.cell(row,1,safe_text(text));cell.alignment=Alignment(vertical='center',wrap_text=True)
            cell.font=Font(name='Calibri',size=12 if heading else 11,bold=heading,color='FFFFFF' if heading else '292633')
            if heading:cell.fill=PatternFill('solid',fgColor='7150A1')
            sheet.row_dimensions[row].height=28 if heading else 34
            return cell
        block(1,str(protein_display(variant.get('protein_snv_alone')) or variant.get('cdna') or 'Variant')+' · '+str(variant.get('sample','')),heading=True)
        block(2,APP_NAME+' · '+__version__+' · GRCh37 / NM_005157.6 / NP_005148.2 · '+datetime.now().isoformat(timespec='minutes'))
        facts=[('Prøve',variant.get('sample')),('Genomisk posisjon',f"{variant.get('chrom')}:{variant.get('pos')} {variant.get('ref')}>{variant.get('alt')}"),
               ('cDNA',variant.get('cdna')),('Protein, enkeltallel',protein_display(variant.get('protein_snv_alone'))),
               ('AF (%)',variant.get('format_af')*100 if variant.get('format_af') is not None else 'Ikke oppgitt'),
               ('ALT / DP',f"{variant.get('alt_ad','–')} / {variant.get('format_dp','–')}"),('FILTER',variant.get('filter'))]
        for row,(name,value) in enumerate(facts,4):
            sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=2)
            sheet.cell(row,1,name).font=Font(bold=True,color='7150A1')
            sheet.merge_cells(start_row=row,start_column=3,end_row=row,end_column=8)
            cell=sheet.cell(row,3,safe_text(value));cell.alignment=Alignment(wrap_text=True,vertical='top')
            sheet.row_dimensions[row].height=24
            if name=='AF (%)':cell.number_format='0.00'
            if name=='Protein, enkeltallel' and flags(variant):cell.comment=Comment(safe_text('; '.join(flags(variant))),APP_NAME)
        # All selected sources remain visible above the image appendix, on the
        # same variant sheet. No overview/data/attachment worksheet is created.
        sources={r['provider']:r for r in lookup_results.get(allele_key(variant),[])}
        row=12
        for entry in variant.get('eln_watchlist',[]):
            block(row,'Variantliste · '+entry['source']+' · '+entry.get('scope',''));row+=2
            if entry.get('recommended_tki'):
                block(row,'TKI oppført i kildetabellen: '+entry['recommended_tki']);row+=2
        block(row,'Databaser',heading=True);row+=1
        for provider in providers:
            result=sources.get(provider)
            sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=2)
            cell=sheet.cell(row,1,PROVIDER_NAMES.get(provider,provider));cell.font=Font(bold=True,color='7150A1')
            sheet.merge_cells(start_row=row,start_column=3,end_row=row,end_column=8)
            value=sheet.cell(row,3,safe_text(source_text(result)));value.alignment=Alignment(wrap_text=True,vertical='top')
            sheet.row_dimensions[row].height=min(110,max(34,18*math.ceil(len(str(value.value))/115)))
            if result:
                value.comment=Comment(safe_text('\n'.join(str(result.get(k) or '') for k in ('checked_at','accession','identity_method'))),APP_NAME)
                url=str(result.get('url') or '')
                if urlsplit(url).scheme=='https':cell.hyperlink=url
            row+=1
        row+=1;block(row,'IGV',heading=True);row+=1
        picture=(igv_images or {}).get(row_key(variant))
        if picture:
            try:row=_image(sheet,picture,row)
            except (OSError,ValueError):block(row,'IGV: Bildet kunne ikke leses.');row+=2
        else:
            block(row,'Ikke tilgjengelig: '+str((igv_errors or {}).get(row_key(variant)) or 'BAM/CRAM og kontrollert referanse kreves.'));row+=2
        for provider in providers:
            result=sources.get(provider);cell=block(row,PROVIDER_NAMES.get(provider,provider),heading=True);row+=1
            text=block(row,source_text(result));sheet.row_dimensions[row].height=64
            if result:
                text.comment=Comment(safe_text('\n'.join(str(result.get(k) or '') for k in ('provider','checked_at','accession','url','identity_method'))),APP_NAME)
                url=str(result.get('url') or '')
                if urlsplit(url).scheme=='https':cell.hyperlink=url
            row+=2
            if result and result.get('status')=='found':
                for key,title in (('transcripts','Transkripter i kilden'),('proteins','Proteiner i kilden'),('reference_hgvs','HGVS i kilden')):
                    values=result.get(key) or []
                    if values:
                        for start in range(0,len(values),4):
                            block(row,title+': '+'; '.join(str(v) for v in values[start:start+4]));sheet.row_dimensions[row].height=64;row+=3
            pictures=screenshot_records(result)
            for picture in pictures:
                caption=str(picture.get('label') or 'Variantresultat')
                captured=picture.get('captured_at') or (result or {}).get('captured_at')
                if captured:caption+=' · '+str(captured)
                block(row,caption);row+=2
                try:row=_image(sheet,picture['path'],row)
                except (OSError,ValueError):block(row,'Bildet kunne ikke leses. Hent bildet på nytt.');row+=2
            if result and result.get('screenshot_error'):block(row,'Bilde: '+str(result['screenshot_error']));row+=2
            elif not pictures:block(row,'Ingen kildebilder hentet.');row+=2
        sheet.freeze_panes='C4';sheet.print_options.horizontalCentered=True
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows='1:2'
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    descriptor,temporary=tempfile.mkstemp(prefix='.DRAGEN-',suffix='.xlsx',dir=path.parent);os.close(descriptor)
    try:book.save(temporary);os.replace(temporary,path)
    finally:
        book.close()
        if Path(temporary).exists():Path(temporary).unlink()

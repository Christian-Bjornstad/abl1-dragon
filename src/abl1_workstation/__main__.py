"""Standalone and future eMolPat-compatible Python entry point."""
from __future__ import annotations
import argparse
import importlib.util
import json
from pathlib import Path
import sys
from . import __version__


def main(argv=None):
    parser=argparse.ArgumentParser(description='ABL1 Dragon')
    parser.add_argument('--project',type=Path,help='Åpne lokalt lagret prosjekt eller analyse-JSON')
    parser.add_argument('--data-dir',type=Path,help='Valgfri godkjent mappe for arbeidsdatabasen')
    parser.add_argument('--smoke-test',type=Path,help='Kjør native selvtest uten WebEngine og lagre kvittering')
    parser.add_argument('--screenshot',type=Path,help='Lagre native vindusgjengivelse under selvtest')
    args=parser.parse_args(argv)
    from PyQt6.QtCore import QCoreApplication, QStandardPaths, Qt, qVersion
    from PyQt6.QtGui import QFont, QFontDatabase, QIcon, QImageReader
    from PyQt6.QtWidgets import QApplication
    QCoreApplication.setAttribute(Qt.ApplicationAttribute.AA_ShareOpenGLContexts)
    if sys.platform=='win32':
        import ctypes
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID('ABL1Dragon.Desktop')
    QImageReader.setAllocationLimit(64)
    app=QApplication(sys.argv[:1]);app.setStyle('Fusion')
    if app.platformName()=='offscreen' and sys.platform=='win32':
        import os
        font_dir=Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'
        for name in ('segoeui.ttf','segoeuib.ttf','seguisb.ttf'):
            font=font_dir/name
            if font.is_file():QFontDatabase.addApplicationFont(str(font))
    app.setFont(QFont('Segoe UI',10))
    app.setOrganizationName('ABL1Dragon');app.setApplicationName('ABL1Dragon');app.setApplicationVersion(__version__)
    app.setApplicationDisplayName('ABL1 Dragon')
    from .domain import ReviewStore, export_xlsx, load_inputs, resource_root, save_project
    icon=resource_root()/'assets/brand/abl1-dragon.ico'
    if icon.is_file():app.setWindowIcon(QIcon(str(icon)))
    from .ui import MainWindow
    data_dir=args.data_dir or Path(QStandardPaths.writableLocation(QStandardPaths.StandardLocation.AppLocalDataLocation))
    store=ReviewStore(data_dir/'vurderinger.sqlite3')
    session=load_inputs([args.project],'GRCh37') if args.project else None
    window=MainWindow(store,browser_enabled=not bool(args.smoke_test),session=session)
    window.show()
    if args.smoke_test:
        app.processEvents()
        window.select_variant('DEMO-A',133748283);window.detail_tabs.setCurrentIndex(1);app.processEvents()
        native_check=window.current_variant and window.current_variant['sample']=='DEMO-A'
        webengine=bool(importlib.util.find_spec('PyQt6.QtWebEngineWidgets'))
        if webengine:
            from PyQt6 import QtWebEngineWidgets  # verify frozen runtime import, no page or browser created
        proof={'application':'ABL1 Dragon','version':__version__,'qt_runtime':qVersion(),
               'frozen':bool(getattr(sys,'frozen',False)),'samples':len(window.session.data['samples']),
               'variants':len(window.session.data['variants']),'qc_rows':len(window.session.data.get('qc',[])),
               'phase_rows':len(window.session.data.get('phase_rows',[])),
               'native_demo_selection':bool(native_check),'synthetic_demo':window.session.data.get('metadata',{}).get('synthetic_demo') is True,'webengine_importable':webengine,
               'webengine_rendering_tested':False,
               'database_provider_count':window.database_panel.providers.count(),
               'edge_profile_launched':False,'database_authentication_tested':False}
        args.smoke_test.parent.mkdir(parents=True,exist_ok=True)
        diagnostic_dir=data_dir/'selvtest';diagnostic_dir.mkdir(parents=True,exist_ok=True)
        project_file=diagnostic_dir/'analyse.abl1.json';save_project(window.session,store,project_file)
        restored=load_inputs([project_file],'GRCh37')
        proof['project_roundtrip']=restored.run_id==window.session.run_id
        workbook_file=diagnostic_dir/'analyse.xlsx';export_xlsx(window.session,store,workbook_file)
        from openpyxl import load_workbook
        workbook=load_workbook(workbook_file,read_only=True)
        proof['excel_variant_rows']=workbook['Varianter'].max_row-1;workbook.close()
        import gzip
        vcf=diagnostic_dir/'NySample.vcf.gz'
        text='##fileformat=VCFv4.2\n#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\tFORMAT\tNySample\nchr9\t133748283\t.\tC\tT\t.\tPASS\t.\tGT:AF:AD:DP\t0/1:0.25:75,25:100\n'
        vcf.write_bytes(gzip.compress(text.encode('utf-8')))
        imported=load_inputs([vcf],'GRCh37');proof['vcf_gzip_import']=len(imported.data['variants'])==1 and not imported.is_pilot and not imported.data['qc']
        if args.screenshot:
            args.screenshot.parent.mkdir(parents=True,exist_ok=True)
            if not window.grab().save(str(args.screenshot)):raise RuntimeError('Skjermbilde kunne ikke lagres')
        args.smoke_test.parent.mkdir(parents=True,exist_ok=True);args.smoke_test.write_text(json.dumps(proof,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        window.close();app.processEvents()
        return 0 if native_check and proof['synthetic_demo'] and proof['project_roundtrip'] and proof['vcf_gzip_import'] and proof['excel_variant_rows']==3 and proof['qc_rows']==0 and proof['phase_rows']==0 else 2
    return app.exec()


if __name__=='__main__':raise SystemExit(main())
